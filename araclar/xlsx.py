import json,csv
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
R=json.load(open('tablo.json'))
EK={'4961':"Varak sayısı (473) tutuyor; ancak Hacıhasanzâde kaydı “hûb hatla … Vakf-ı Hatîbzâde” diyor, Fenârî kaydı “bi-hattı … Musannifek” diyor. Veritabanında bu iki kayıt birbirine ilişkili değil. Kontrol edilmeli.",
    '110':"Hacıhasanzâde kaydında sıra numarası yazılmamış (63a'daki ilk kayıt)."}
H=["Sıra (listedeki)","Fatih nr.","Fenârî varak [sıra]","Fenârî kataloğu – tam metin","Hacıhasanzâde varak [sıra]","Hacıhasanzâde kataloğu – tam metin","Veritabanı nüshası","Veritabanı kayıt no (Fenârî / Hacıhasanzâde)","Not"]
rows=[]
for o in R:
    note=o['not']
    if o['fatih']=='110': note=note.replace("Hacıhasanzâde ref. listede ['63a', 1], veritabanında 63a; ","")
    note="; ".join(x for x in [EK.get(o['fatih']),note] if x)
    F,Hh=o['F'],o['H']
    rows.append([o['sira'],"Fatih "+o['fatih'],F['ref'] if F else "—",F['text'] if F else "",Hh['ref'] if Hh else "—",Hh['text'] if Hh else "",
                 o['yontem'].replace('nüsha kaydı (','').rstrip(')').replace('nüsha kaydı yok','yok'),
                 f"{'#'+str(F['id']) if F else '—'} / {'#'+str(Hh['id']) if Hh else '—'}",note])
wb=Workbook(); ws=wb.active; ws.title="Yazma–katalog"
ws.append(H)
for r in rows: ws.append(r)
W=[10,13,12,70,13,70,14,18,55]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
hf=PatternFill("solid",fgColor="DDE4EE"); warn=PatternFill("solid",fgColor="FCE8C8")
for c in ws[1]: c.font=Font(bold=True); c.fill=hf; c.alignment=Alignment(wrap_text=True,vertical="center")
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    if any(k in (row[8].value or '') for k in ("ÇELİŞKİ","YOK","Kontrol","doğrusu")):
        for c in row: c.fill=warn
ws.freeze_panes="C2"; ws.auto_filter.ref=ws.dimensions
wb.save('/home/user/Denemeler/yazma_katalog_eslesme.xlsx')
with open('/home/user/Denemeler/yazma_katalog_eslesme.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(H); w.writerows(rows)
print(len(rows), sum(1 for r in rows if r[3]), sum(1 for r in rows if r[5]))
